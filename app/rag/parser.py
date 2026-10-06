"""Document parsing: PDF (text and scanned), DOCX, TXT/Markdown and images -> a structured block list.

The parser's job is to keep structure (headings, clauses, lists, tables, page numbers) so the chunker can
preserve the document hierarchy. It never invents text; OCR output carries a confidence.
"""
from __future__ import annotations

import io
import re
import statistics
from collections import Counter
from dataclasses import dataclass, field
from typing import Optional

from app.core.config import get_settings
from app.core.errors import OCRUnavailable, ParseError
from app.core.logging import get_logger
from app.rag import ocr

log = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}


@dataclass
class Block:
    kind: str  # heading | paragraph | list_item | clause | table
    text: str
    level: int = 0  # heading depth (1 = top)
    number: str = ""  # section/clause number as written, e.g. "7.2.1"
    page: Optional[int] = None
    rows: Optional[list[list[str]]] = None
    ocr: bool = False
    confidence: float = 1.0


@dataclass
class ParsedDocument:
    blocks: list[Block]
    page_count: int = 0
    ocr_used: bool = False
    warnings: list[str] = field(default_factory=list)
    mean_confidence: float = 1.0

    @property
    def char_count(self) -> int:
        return sum(len(b.text) for b in self.blocks)


# ------------------------------------------------------------------------------------------------
# Shared structure detection
# ------------------------------------------------------------------------------------------------
_NUM_RE = re.compile(r"^(\d+(?:\.\d+){0,5})[.)]?\s+(\S.*)$")
_KEYWORD_HEAD_RE = re.compile(r"^(?:section|chapter|article|part|annexure|annex|appendix|schedule)\s+([A-Za-z0-9.]+)\b\s*[:.\-–\u2014]?\s*(.*)$", re.IGNORECASE)
_MD_HEAD_RE = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")
_BULLET_RE = re.compile(r"^\s*(?:[•●▪‣◦*\-–]|\(?[a-zA-Z]{1,2}[).]|\(?[ivxIVX]{1,4}[).]|\d{1,2}[)])\s+(\S.*)$")
_TOC_RE = re.compile(r"\.{4,}\s*\d+\s*$")
_ALLCAPS_RE = re.compile(r"^[A-Z0-9][A-Z0-9 &,/()\-:'–]{3,78}$")
_PAGE_NUM_RE = re.compile(r"^\s*(?:page\s+)?\d+(?:\s*(?:of|/)\s*\d+)?\s*$", re.IGNORECASE)


def classify_text(text: str, hint_bold: bool = False, hint_large: bool = False) -> Block:
    """Classify one paragraph-sized piece of text into a Block (no page information)."""
    t = re.sub(r"\s+", " ", text).strip()
    m = _MD_HEAD_RE.match(t)
    if m:
        title = m.group(2).strip()
        n = _NUM_RE.match(title)
        return Block("heading", n.group(2) if n else title, level=len(m.group(1)), number=n.group(1) if n else "")
    n = _NUM_RE.match(t)
    if n:
        number, rest = n.group(1), n.group(2).strip()
        looks_like_title = len(rest) <= 100 and not rest.endswith((".", ";", ",", ":")) and len(rest.split()) <= 14
        # "2026 batch ..." or "1 Introduction" -- a lone year-like number is not a section number
        if re.fullmatch(r"(19|20)\d{2}", number):
            looks_like_title = False
        if looks_like_title and (hint_bold or hint_large or rest[:1].isupper()):
            return Block("heading", rest, level=number.count(".") + 1, number=number)
        if "." in number or len(rest) > 100 or not looks_like_title:
            if not re.fullmatch(r"(19|20)\d{2}", number):
                return Block("clause", rest, number=number)
    k = _KEYWORD_HEAD_RE.match(t)
    if k and len(t) <= 110 and not t.endswith("."):
        return Block("heading", (k.group(2) or t).strip() or t, level=1, number=t.split()[0].capitalize() + " " + k.group(1))
    if (hint_bold or hint_large) and len(t) <= 110 and not t.endswith((".", ";", ",")):
        return Block("heading", t, level=2 if not hint_large else 1)
    if _ALLCAPS_RE.match(t) and sum(c.isalpha() for c in t) >= 4 and len(t.split()) <= 10:
        return Block("heading", t.title(), level=1)
    b = _BULLET_RE.match(t)
    if b:
        return Block("list_item", t)
    return Block("paragraph", t)


def _is_pipe_table(lines: list[str]) -> bool:
    return len(lines) >= 2 and all(l.count("|") >= 1 for l in lines)


def _parse_pipe_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for l in lines:
        cells = [c.strip() for c in l.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue  # markdown separator row
        rows.append(cells)
    return rows


# ------------------------------------------------------------------------------------------------
# TXT / Markdown
# ------------------------------------------------------------------------------------------------
def parse_text(data: bytes) -> ParsedDocument:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    blocks: list[Block] = []
    page = 1
    para: list[str] = []

    def flush() -> None:
        nonlocal para
        if not para:
            return
        lines, para = para, []
        current: list[str] = []
        table_lines: list[str] = []

        def emit_current() -> None:
            nonlocal current
            if current:
                b = classify_text(" ".join(current))
                b.page = page
                blocks.append(b)
                current = []

        def emit_table() -> None:
            nonlocal table_lines
            if not table_lines:
                return
            if _is_pipe_table(table_lines):
                rows = _parse_pipe_table(table_lines)
                blocks.append(Block("table", " ".join(" ".join(r) for r in rows), rows=rows, page=page))
            else:  # a single line containing a pipe is ordinary text
                for l in table_lines:
                    b = classify_text(l)
                    b.page = page
                    blocks.append(b)
            table_lines = []

        for line in lines:
            stripped = line.strip()
            if stripped.count("|") >= 2 or (table_lines and "|" in stripped):
                emit_current()
                table_lines.append(stripped)
                continue
            emit_table()
            starts_item = bool(_BULLET_RE.match(line) or _MD_HEAD_RE.match(line) or _NUM_RE.match(line) or _KEYWORD_HEAD_RE.match(line))
            if starts_item:
                emit_current()
                head = classify_text(stripped)
                if head.kind == "heading":
                    head.page = page
                    blocks.append(head)  # a heading is always its own block, never merged with the lines below it
                else:
                    current = [stripped]
            else:
                current.append(stripped)
        emit_current()
        emit_table()

    for raw in text.splitlines():
        if "\f" in raw:
            flush()
            page += raw.count("\f")
            raw = raw.replace("\f", "")
        if not raw.strip():
            flush()
            continue
        if _TOC_RE.search(raw):
            continue
        para.append(raw)
    flush()
    if not blocks:
        raise ParseError("the text file contains no readable text")
    return ParsedDocument(blocks=blocks, page_count=page)


# ------------------------------------------------------------------------------------------------
# DOCX
# ------------------------------------------------------------------------------------------------
def parse_docx(data: bytes) -> ParsedDocument:
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise ParseError("python-docx is not installed") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"the DOCX file is corrupted or unreadable: {exc}") from exc
    blocks: list[Block] = []
    warnings: list[str] = []
    page = 1
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, document)
            if 'w:br w:type="page"' in child.xml or "w:lastRenderedPageBreak" in child.xml:
                page += child.xml.count('w:type="page"') + child.xml.count("lastRenderedPageBreak")
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"(?i)^heading\s*(\d)", style)
            if m:
                n = _NUM_RE.match(text)
                blocks.append(Block("heading", n.group(2) if n else text, level=int(m.group(1)), number=n.group(1) if n else "", page=page))
            elif style.lower() == "title":
                blocks.append(Block("heading", text, level=1, page=page))
            elif "list" in style.lower():
                blocks.append(Block("list_item", text, page=page))
            else:
                bold = bool(p.runs) and all((r.bold or False) for r in p.runs if r.text.strip())
                b = classify_text(text, hint_bold=bold)
                b.page = page
                blocks.append(b)
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for r in table.rows:
                cells = []
                prev = None
                for c in r.cells:
                    if c._tc is prev:  # merged cell repeated by python-docx
                        continue
                    prev = c._tc
                    cells.append(re.sub(r"\s+", " ", c.text).strip())
                if any(cells):
                    rows.append(cells)
            if len(rows) >= 1:
                blocks.append(Block("table", " ".join(" ".join(r) for r in rows), rows=rows, page=page))
    if not blocks:
        raise ParseError("the DOCX file contains no readable text")
    return ParsedDocument(blocks=blocks, page_count=page, warnings=warnings)


# ------------------------------------------------------------------------------------------------
# Images (OCR)
# ------------------------------------------------------------------------------------------------
def _blocks_from_ocr(page_no: int, page: "ocr.OCRPage") -> list[Block]:
    if not page.lines:
        return []
    heights = [l.height for l in page.lines]
    med = statistics.median(heights) if heights else 0
    gaps = [b.top - (a.top + a.height) for a, b in zip(page.lines, page.lines[1:])]
    gap_med = statistics.median(gaps) if gaps else 0
    blocks: list[Block] = []
    buf: list[ocr.OCRLine] = []

    def flush() -> None:
        nonlocal buf
        if not buf:
            return
        text = " ".join(l.text for l in buf)
        large = max(l.height for l in buf) >= 1.3 * med and len(buf) == 1
        b = classify_text(text, hint_large=large)
        b.page = page_no
        b.ocr = True
        b.confidence = sum(l.confidence for l in buf) / len(buf)
        blocks.append(b)
        buf = []

    for i, line in enumerate(page.lines):
        text = line.text.strip()
        starts_new = bool(_NUM_RE.match(text) or _BULLET_RE.match(text) or _KEYWORD_HEAD_RE.match(text) or _MD_HEAD_RE.match(text))
        if buf and (starts_new or (i > 0 and gaps[i - 1] > max(1.0, gap_med) * 1.8 + 4)):
            flush()
        buf.append(line)
    flush()
    return blocks


def parse_image(data: bytes) -> ParsedDocument:
    page = ocr.ocr_image_bytes(data)
    blocks = _blocks_from_ocr(1, page)
    if not blocks:
        raise ParseError("OCR found no text in the image")
    return ParsedDocument(blocks=blocks, page_count=1, ocr_used=True, mean_confidence=page.mean_confidence)


# ------------------------------------------------------------------------------------------------
# PDF
# ------------------------------------------------------------------------------------------------
def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        import pymupdf
    except ImportError as exc:  # pragma: no cover
        raise ParseError("PyMuPDF is not installed") from exc
    settings = get_settings()
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ParseError(f"the PDF is corrupted or unreadable: {exc}") from exc
    if pdf.is_encrypted and not pdf.authenticate(""):
        raise ParseError("the PDF is password protected")

    warnings: list[str] = []
    raw_pages: list[dict] = []
    sizes: list[tuple[float, int]] = []
    for pno in range(pdf.page_count):
        page = pdf.load_page(pno)
        height = float(page.rect.height)
        table_boxes: list[tuple[float, float, float, float]] = []
        table_blocks: list[tuple[float, Block]] = []
        try:
            for t in page.find_tables().tables:
                rows = [[re.sub(r"\s+", " ", (c or "")).strip() for c in row] for row in t.extract()]
                rows = [r for r in rows if any(r)]
                if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
                    table_boxes.append(tuple(t.bbox))  # type: ignore[arg-type]
                    table_blocks.append((t.bbox[1], Block("table", " ".join(" ".join(r) for r in rows), rows=rows, page=pno + 1)))
        except Exception as exc:  # table detection is best-effort
            warnings.append(f"page {pno + 1}: table detection failed ({type(exc).__name__})")
        text_blocks: list[dict] = []
        for blk in page.get_text("dict").get("blocks", []):
            if blk.get("type") != 0:
                continue
            x0, y0, x1, y1 = blk["bbox"]
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            if any(tb[0] <= cx <= tb[2] and tb[1] <= cy <= tb[3] for tb in table_boxes):
                continue
            lines = []
            blk_sizes, bold_chars, chars = [], 0, 0
            for ln in blk.get("lines", []):
                spans = ln.get("spans", [])
                txt = "".join(s.get("text", "") for s in spans).strip()
                if not txt:
                    continue
                lines.append(txt)
                for s in spans:
                    n = len(s.get("text", "").strip())
                    if n:
                        blk_sizes.append(round(float(s.get("size", 0)), 1))
                        sizes.append((round(float(s.get("size", 0)), 1), n))
                        chars += n
                        if (int(s.get("flags", 0)) & 16) or "bold" in str(s.get("font", "")).lower():
                            bold_chars += n
            if not lines:
                continue
            text_blocks.append({"y0": y0, "y1": y1, "lines": lines, "size": max(blk_sizes) if blk_sizes else 0,
                                "bold": chars > 0 and bold_chars / chars > 0.6})
        char_total = sum(len(" ".join(b["lines"])) for b in text_blocks)
        raw_pages.append({"pno": pno + 1, "height": height, "text_blocks": text_blocks, "table_blocks": table_blocks,
                          "chars": char_total, "page": page})

    body_size = 0.0
    if sizes:
        counter: Counter = Counter()
        for sz, n in sizes:
            counter[sz] += n
        body_size = counter.most_common(1)[0][0]

    # running headers/footers: short text repeating at the top/bottom of >= 50% of pages
    repeated = _repeated_margin_text(raw_pages)

    blocks: list[Block] = []
    ocr_used = False
    confidences: list[float] = []
    heading_sizes = sorted({round(b["size"], 1) for p in raw_pages for b in p["text_blocks"] if body_size and b["size"] >= body_size * 1.12}, reverse=True)
    for rp in raw_pages:
        scanned = rp["chars"] < settings.ocr_min_chars_per_page
        if scanned:
            try:
                png = rp["page"].get_pixmap(dpi=200).tobytes("png")
                page_ocr = ocr.ocr_image_bytes(png)
                ocr_blocks = _blocks_from_ocr(rp["pno"], page_ocr)
                if ocr_blocks:
                    ocr_used = True
                    confidences.append(page_ocr.mean_confidence)
                    blocks.extend(ocr_blocks)
                    if page_ocr.mean_confidence < 0.6:
                        warnings.append(f"page {rp['pno']}: low OCR confidence ({page_ocr.mean_confidence:.2f}); text may contain errors")
                    continue
                warnings.append(f"page {rp['pno']}: no text found, even with OCR")
            except OCRUnavailable as exc:
                warnings.append(f"page {rp['pno']}: scanned page skipped, OCR unavailable ({exc})")
            continue
        items: list[tuple[float, Block]] = list(rp["table_blocks"])
        for tb in rp["text_blocks"]:
            joined = " ".join(tb["lines"])
            norm = _margin_key(joined)
            near_edge = tb["y0"] < rp["height"] * 0.08 or tb["y1"] > rp["height"] * 0.92
            if near_edge and (norm in repeated or _PAGE_NUM_RE.match(joined)):
                continue
            if _TOC_RE.search(joined):
                continue
            large = bool(body_size) and tb["size"] >= body_size * 1.12
            # multi-line blocks that are really list items / numbered clauses: split per line start
            pieces = _split_block_lines(tb["lines"])
            for piece in pieces:
                b = classify_text(piece, hint_bold=tb["bold"] and len(pieces) == 1, hint_large=large and len(pieces) == 1)
                if b.kind == "heading" and large:
                    b.level = min(b.level or 1, 1 + heading_sizes.index(round(tb["size"], 1))) if round(tb["size"], 1) in heading_sizes else b.level
                b.page = rp["pno"]
                items.append((tb["y0"], b))
        items.sort(key=lambda it: it[0])
        blocks.extend(b for _, b in items)
    blocks = _normalise_heading_levels(blocks)
    if not blocks:
        raise ParseError("no text could be extracted from the PDF (it may be scanned and OCR may be unavailable)")
    mean_conf = sum(confidences) / len(confidences) if confidences else 1.0
    return ParsedDocument(blocks=blocks, page_count=pdf.page_count, ocr_used=ocr_used, warnings=warnings, mean_confidence=mean_conf)


def _margin_key(text: str) -> str:
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", text.strip().lower()))


def _repeated_margin_text(raw_pages: list[dict]) -> set[str]:
    if len(raw_pages) < 2:
        return set()
    counts: Counter = Counter()
    for rp in raw_pages:
        seen = set()
        for tb in rp["text_blocks"]:
            if tb["y0"] < rp["height"] * 0.08 or tb["y1"] > rp["height"] * 0.92:
                key = _margin_key(" ".join(tb["lines"]))
                if key and len(key) < 140:
                    seen.add(key)
        counts.update(seen)
    need = max(2, int(len(raw_pages) * 0.5))
    return {k for k, v in counts.items() if v >= need}


def _split_block_lines(lines: list[str]) -> list[str]:
    """A PDF text block can hold several list items or numbered clauses; split where a new item starts."""
    pieces: list[str] = []
    for line in lines:
        starts = bool(_BULLET_RE.match(line) or _NUM_RE.match(line) or _KEYWORD_HEAD_RE.match(line))
        if starts or not pieces:
            pieces.append(line.strip())
        else:
            pieces[-1] = pieces[-1] + " " + line.strip()
    return pieces


def _normalise_heading_levels(blocks: list[Block]) -> list[Block]:
    """Numbered headings use their depth; un-numbered ones keep their detected level."""
    for b in blocks:
        if b.kind == "heading" and b.number and re.fullmatch(r"\d+(?:\.\d+)*", b.number):
            b.level = b.number.count(".") + 1
    return blocks


# ------------------------------------------------------------------------------------------------
# Entry point
# ------------------------------------------------------------------------------------------------
def parse_document(data: bytes, filename: str) -> ParsedDocument:
    name = (filename or "").lower()
    ext = "." + name.rsplit(".", 1)[-1] if "." in name else ""
    if ext == ".doc":
        raise ParseError("legacy .doc files are not supported; please save the document as .docx or PDF")
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError(f"unsupported file type '{ext or 'unknown'}'; supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    if not data:
        raise ParseError("the uploaded file is empty")
    if ext == ".pdf":
        if not data.lstrip().startswith(b"%PDF"):
            raise ParseError("the file is not a valid PDF (missing %PDF header)")
        return parse_pdf(data)
    if ext == ".docx":
        return parse_docx(data)
    if ext in (".txt", ".md", ".markdown"):
        return parse_text(data)
    return parse_image(data)

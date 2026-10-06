"""Render DocSpec objects to PDF, DOCX, Markdown and an image-only scanned PDF."""
from __future__ import annotations

import io
import textwrap
from pathlib import Path

from app.synthetic.corpus import INSTITUTE, DocSpec


def render_pdf(spec: DocSpec) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=spec.title, leftMargin=56, rightMargin=56, topMargin=64, bottomMargin=64)
    ss = getSampleStyleSheet()
    h0 = ParagraphStyle("H0", parent=ss["Title"], fontSize=20, spaceAfter=6)
    h1 = ParagraphStyle("H1", parent=ss["Heading2"], fontSize=14, spaceBefore=10)
    body = ParagraphStyle("B", parent=ss["BodyText"], fontSize=10.5, leading=15)
    small = ParagraphStyle("S", parent=ss["BodyText"], fontSize=9, textColor=colors.grey)

    def header_footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillGray(0.45)
        canvas.drawString(56, A4[1] - 36, f"{INSTITUTE} | {spec.doc_id} | version {spec.version}")
        canvas.drawRightString(A4[0] - 56, 32, f"Page {d.page}")
        canvas.restoreState()

    story = [Paragraph(spec.title, h0), Paragraph(f"Issued by: {spec.issuer}. Effective from {spec.effective_from}. Version {spec.version}.", small), Spacer(1, 8)]
    if spec.intro:
        story.append(Paragraph(spec.intro, body))
    for i, sec in enumerate(spec.sections):
        story.append(Paragraph(f"{sec.number} {sec.title}", h1))
        for p in sec.paragraphs:
            story.append(Paragraph(p, body))
        if sec.table:
            if sec.table_caption:
                story.append(Paragraph(sec.table_caption, small))
            t = Table(sec.table, hAlign="LEFT")
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
            story.append(t)
        if i == 2 and len(spec.sections) > 4:
            story.append(PageBreak())  # a second page so that page numbers in citations are meaningful
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    return buf.getvalue()


def render_docx(spec: DocSpec) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(spec.title, level=1)
    d.add_paragraph(f"Issued by: {spec.issuer}. Effective from {spec.effective_from}. Version {spec.version}.")
    if spec.intro:
        d.add_paragraph(spec.intro)
    for sec in spec.sections:
        d.add_heading(f"{sec.number} {sec.title}", level=2)
        for p in sec.paragraphs:
            d.add_paragraph(p)
        if sec.table:
            t = d.add_table(rows=len(sec.table), cols=len(sec.table[0]))
            t.style = "Table Grid"
            for r, row in enumerate(sec.table):
                for c, val in enumerate(row):
                    t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_md(spec: DocSpec) -> bytes:
    lines = [f"# {spec.title}", "", f"Issued by: {spec.issuer}. Effective from {spec.effective_from}. Version {spec.version}.", ""]
    if spec.intro:
        lines += [spec.intro, ""]
    for sec in spec.sections:
        lines += [f"## {sec.number} {sec.title}", ""]
        lines += [p + "\n" for p in sec.paragraphs]
        if sec.table:
            lines += ["| " + " | ".join(sec.table[0]) + " |", "|" + "---|" * len(sec.table[0])]
            lines += ["| " + " | ".join(r) + " |" for r in sec.table[1:]]
            lines.append("")
    return "\n".join(lines).encode("utf-8")


def render_scanned(spec: DocSpec) -> bytes:
    """Image-only PDF: forces the OCR path (no text layer)."""
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("L", (1240, 1754), 255)
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34)
        big = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 46)
    except OSError:
        font = big = ImageFont.load_default()
    y = 110
    draw.text((100, y), spec.title, fill=0, font=big)
    y += 100
    for sec in spec.sections:
        draw.text((100, y), f"{sec.number} {sec.title}", fill=0, font=big)
        y += 80
        for p in sec.paragraphs:
            for line in textwrap.wrap(p, 62):
                draw.text((100, y), line, fill=0, font=font)
                y += 52
            y += 14
    buf = io.BytesIO()
    img.save(buf, format="PDF", resolution=150)
    return buf.getvalue()


RENDERERS = {"pdf": render_pdf, "docx": render_docx, "md": render_md, "scanned": render_scanned}


def render(spec: DocSpec) -> bytes:
    return RENDERERS[spec.fmt](spec)


def write_all(specs: list[DocSpec], out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for s in specs:
        p = out_dir / s.filename
        p.write_bytes(render(s))
        paths.append(p)
    return paths
